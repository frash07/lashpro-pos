import csv
from collections import Counter
from decimal import Decimal, InvalidOperation
import io
import re
import secrets
import unicodedata
import zipfile

from django.core.exceptions import ValidationError
from django.db import IntegrityError, transaction
from django.db.models.functions import Lower
from openpyxl import load_workbook

from accounts.models import User
from audit.models import AuditEvent
from audit.services import record_audit_event
from inventory.services import receive_stock
from .models import Product, SKUReservation

MAX_UPLOAD_BYTES = 5 * 1024 * 1024
MAX_UNCOMPRESSED_XLSX_BYTES = 25 * 1024 * 1024
MAX_IMPORT_ROWS = 2000
MAX_SKU_LENGTH = 20
PRICE_MAX = Decimal("99999999.99")
QUANTITY_MAX = Decimal("999999999.999")
SKU_ALPHABET = "ABCDEFGHJKLMNPQRSTUVWXYZ23456789"
SKU_PATTERN = re.compile(r"^[A-HJ-NP-Z0-9]+(?:-[A-HJ-NP-Z0-9]+)*$")
FORMULA_CELL = object()

ALIASES = {
    "product": "name",
    "product_name": "name",
    "item_name": "name",
    "sku_code": "sku",
    "cost": "cost_price",
    "unit_cost": "cost_price",
    "purchase_price": "cost_price",
    "price": "selling_price",
    "sale_price": "selling_price",
    "retail_price": "selling_price",
    "opening_quantity": "opening_stock",
    "on_hand": "opening_stock",
    "stock": "opening_stock",
    "reorder_point": "reorder_level",
    "minimum_stock": "reorder_level",
}
RECOGNIZED_COLUMNS = {"name", "sku", "cost_price", "selling_price", "opening_stock", "reorder_level"}
REQUIRED_COLUMNS = {"name", "cost_price", "selling_price"}


def _clean_header(value):
    if value is None or value is FORMULA_CELL:
        return ""
    key = re.sub(r"[^a-z0-9]+", "_", str(value).strip().casefold()).strip("_")
    return ALIASES.get(key, key)


def _xlsx_rows(raw):
    try:
        with zipfile.ZipFile(io.BytesIO(raw)) as archive:
            members = archive.infolist()
            if len(members) > 1000 or sum(item.file_size for item in members) > MAX_UNCOMPRESSED_XLSX_BYTES:
                raise ValidationError("This Excel workbook is too large to import safely.")
        workbook = load_workbook(io.BytesIO(raw), read_only=True, data_only=False)
    except ValidationError:
        raise
    except Exception as exc:
        raise ValidationError("The Excel file could not be read. Upload a valid .xlsx workbook.") from exc

    try:
        sheet = workbook.active
        if sheet is None:
            raise ValidationError("The Excel workbook does not contain a worksheet.")
        if sheet.max_row and sheet.max_row > MAX_IMPORT_ROWS + 1:
            raise ValidationError(f"Import files may contain at most {MAX_IMPORT_ROWS} product rows. Remove extra formatted rows from the worksheet.")
        rows = []
        for row_number, cells in enumerate(sheet.iter_rows(), start=1):
            values = [FORMULA_CELL if cell.data_type == "f" else cell.value for cell in cells]
            rows.append((row_number, values))
            if len(rows) > MAX_IMPORT_ROWS + 1:
                raise ValidationError(f"Import files may contain at most {MAX_IMPORT_ROWS} product rows.")
        return rows
    finally:
        workbook.close()


def _uploaded_rows(uploaded_file):
    extension = uploaded_file.name.rsplit(".", 1)[-1].casefold() if "." in uploaded_file.name else ""
    if extension not in {"csv", "xlsx"}:
        raise ValidationError("Upload a CSV or .xlsx file.")
    raw = uploaded_file.read(MAX_UPLOAD_BYTES + 1)
    if len(raw) > MAX_UPLOAD_BYTES:
        raise ValidationError("The file is too large. Maximum upload size is 5 MB.")

    if extension == "xlsx":
        rows = _xlsx_rows(raw)
    else:
        try:
            text = raw.decode("utf-8-sig")
            reader = csv.reader(io.StringIO(text, newline=""), strict=True)
            rows = [(row_number, values) for row_number, values in enumerate(reader, start=1)]
        except (UnicodeDecodeError, csv.Error) as exc:
            raise ValidationError("The CSV must be valid UTF-8 with properly quoted fields.") from exc
        if len(rows) > MAX_IMPORT_ROWS + 1:
            raise ValidationError(f"Import files may contain at most {MAX_IMPORT_ROWS} product rows.")

    if not rows or not any(value not in (None, "") for value in rows[0][1]):
        raise ValidationError("The first row must contain column headings.")
    return rows


def _column_indexes(header_row):
    indexes = {}
    for index, value in enumerate(header_row):
        key = _clean_header(value)
        if key not in RECOGNIZED_COLUMNS:
            continue
        if key in indexes:
            raise ValidationError(f"The heading '{key}' appears more than once.")
        indexes[key] = index
    missing = sorted(REQUIRED_COLUMNS - indexes.keys())
    if missing:
        raise ValidationError("Required columns are missing: " + ", ".join(missing) + ".")
    return indexes


def _decimal_value(value, *, label, row_number, places, maximum, default=None):
    if value is None or (isinstance(value, str) and not value.strip()):
        if default is not None:
            return default
        raise ValidationError(f"Row {row_number}: {label} is required.")
    if isinstance(value, bool):
        raise ValidationError(f"Row {row_number}: {label} must be a number.")
    try:
        amount = Decimal(str(value).strip())
        unit = Decimal(1).scaleb(-places)
        rounded = amount.quantize(unit)
    except (InvalidOperation, TypeError, ValueError):
        raise ValidationError(f"Row {row_number}: enter a valid {label} without currency symbols or thousands separators.") from None
    if not amount.is_finite() or amount != rounded or amount < 0 or amount > maximum:
        raise ValidationError(f"Row {row_number}: {label} must be non-negative, within the supported range, and use at most {places} decimal places.")
    return rounded


def _validated_sku(value, row_number):
    sku = str(value).strip()
    if not sku:
        return ""
    if len(sku) > MAX_SKU_LENGTH or not SKU_PATTERN.fullmatch(sku) or "I" in sku or "O" in sku:
        raise ValidationError(
            f"Row {row_number}: SKU must use uppercase A–Z or 0–9 with optional dashes, exclude I and O, and be no longer than 20 characters."
        )
    return sku


def _row_value(values, indexes, name):
    index = indexes.get(name)
    return values[index] if index is not None and index < len(values) else None


def _normalize_row(row_number, values, indexes):
    errors = []
    if any(value is FORMULA_CELL for value in values):
        errors.append(f"Row {row_number}: formulas are not accepted; replace them with values.")
        values = [None if value is FORMULA_CELL else value for value in values]
    name_value = _row_value(values, indexes, "name")
    name = "" if name_value is None else str(name_value).strip()
    if not name:
        errors.append(f"Row {row_number}: product name is required.")
    elif len(name) > 100:
        errors.append(f"Row {row_number}: product name must be 100 characters or fewer.")

    data = {
        "row_number": row_number,
        "name": name,
        "sku": "",
        "sku_generated": False,
        "cost_price": "",
        "selling_price": "",
        "opening_stock": "0.000",
        "reorder_level": "0.000",
        "errors": errors,
    }
    raw_sku = _row_value(values, indexes, "sku")
    try:
        data["sku"] = _validated_sku(raw_sku, row_number)
        data["sku_generated"] = not data["sku"]
    except ValidationError as exc:
        errors.extend(exc.messages)

    numeric_fields = (
        ("cost_price", "cost price", 2, PRICE_MAX, None),
        ("selling_price", "selling price", 2, PRICE_MAX, None),
        ("opening_stock", "opening stock", 3, QUANTITY_MAX, Decimal("0.000")),
        ("reorder_level", "reorder level", 3, QUANTITY_MAX, Decimal("0.000")),
    )
    for field, label, places, maximum, default in numeric_fields:
        try:
            value = _decimal_value(
                _row_value(values, indexes, field), label=label, row_number=row_number,
                places=places, maximum=maximum, default=default,
            )
            data[field] = format(value, f".{places}f")
        except ValidationError as exc:
            errors.extend(exc.messages)
    return data


def _product_prefix(name):
    normalized = unicodedata.normalize("NFKD", name.upper()).encode("ascii", "ignore").decode("ascii")
    words = re.findall(r"[A-Z]+", normalized)
    initials = [word[0] for word in words if word[0] not in "IO"]
    prefix = initials[:3]
    if len(prefix) < 3:
        for character in "".join(words):
            if character not in "IO" and character not in prefix:
                prefix.append(character)
            if len(prefix) == 3:
                break
    return "".join(prefix) if len(prefix) == 3 else "PRD"


def _generate_sku(name, unavailable):
    prefix = _product_prefix(name)
    for _ in range(100):
        suffix = "".join(secrets.choice(SKU_ALPHABET) for _ in range(5))
        candidate = f"{prefix}-{suffix}"
        if candidate not in unavailable:
            unavailable.add(candidate)
            return candidate
    raise ValidationError("A unique SKU could not be generated. Please preview the import again.")


def _known_skus():
    return {sku.upper() for sku in SKUReservation.objects.values_list("sku", flat=True)}


def preview_product_import(uploaded_file):
    file_rows = _uploaded_rows(uploaded_file)
    indexes = _column_indexes(file_rows[0][1])
    rows = []
    for row_number, values in file_rows[1:]:
        if all(value in (None, "") for value in values):
            continue
        rows.append(_normalize_row(row_number, values, indexes))
    if not rows:
        raise ValidationError("The file does not contain any product rows.")

    used = _known_skus()
    counts = Counter(row["sku"].upper() for row in rows if row["sku"])
    for row in rows:
        sku = row["sku"].upper()
        if not sku:
            continue
        if counts[sku] > 1:
            row["errors"].append(f"Row {row['row_number']}: SKU '{sku}' is duplicated in this file.")
        elif sku in used:
            row["errors"].append(f"Row {row['row_number']}: SKU '{sku}' has already been issued to a product.")
    used.update(row["sku"].upper() for row in rows if row["sku"])
    for row in rows:
        if row["sku_generated"] and row["name"]:
            row["sku"] = _generate_sku(row["name"], used)
    return rows


def _normalize_confirmation_rows(rows):
    if not isinstance(rows, list) or not rows or len(rows) > MAX_IMPORT_ROWS:
        raise ValidationError("The import preview is missing or has too many rows. Upload the file and preview it again.")
    normalized = []
    for item in rows:
        if not isinstance(item, dict):
            raise ValidationError("The import preview is invalid. Upload the file and preview it again.")
        row_number = item.get("row_number")
        try:
            row_number = int(row_number)
        except (TypeError, ValueError):
            row_number = 0
        values = {
            "row_number": row_number,
            "name": str(item.get("name", "")).strip(),
            "sku": "",
            "sku_generated": bool(item.get("sku_generated")),
            "cost_price": "",
            "selling_price": "",
            "opening_stock": "0.000",
            "reorder_level": "0.000",
            "errors": [],
        }
        if not values["name"] or len(values["name"]) > 100:
            values["errors"].append(f"Row {row_number}: product name is invalid.")
        try:
            values["sku"] = _validated_sku(item.get("sku", ""), row_number)
            if not values["sku"]:
                values["errors"].append(f"Row {row_number}: the preview has no SKU.")
        except ValidationError as exc:
            values["errors"].extend(exc.messages)
        for field, label, places, maximum in (
            ("cost_price", "cost price", 2, PRICE_MAX),
            ("selling_price", "selling price", 2, PRICE_MAX),
            ("opening_stock", "opening stock", 3, QUANTITY_MAX),
            ("reorder_level", "reorder level", 3, QUANTITY_MAX),
        ):
            try:
                value = _decimal_value(
                    item.get(field), label=label, row_number=row_number,
                    places=places, maximum=maximum,
                )
                values[field] = format(value, f".{places}f")
            except ValidationError as exc:
                values["errors"].extend(exc.messages)
        normalized.append(values)

    counts = Counter(row["sku"].upper() for row in normalized if row["sku"])
    candidates = list(counts)
    existing = set(
        SKUReservation.objects.annotate(normalized_sku=Lower("sku"))
        .filter(normalized_sku__in=[sku.lower() for sku in candidates])
        .values_list("normalized_sku", flat=True)
    )
    for row in normalized:
        sku = row["sku"].upper()
        if counts[sku] > 1:
            row["errors"].append(f"Row {row['row_number']}: SKU '{sku}' is duplicated in this file.")
        elif sku.lower() in existing:
            row["errors"].append(f"Row {row['row_number']}: SKU '{sku}' has already been issued to a product.")
    errors = [message for row in normalized for message in row["errors"]]
    if errors:
        raise ValidationError(errors)
    return normalized


def _is_manager(actor):
    return bool(
        getattr(actor, "is_active", False)
        and (actor.is_superuser or getattr(actor, "role", None) in {User.Role.OWNER, User.Role.MANAGER})
    )


def _safe_filename(filename):
    name = str(filename or "products upload")
    name = name.replace("\\", "/").rsplit("/", 1)[-1].strip()
    return name[:120] or "products upload"


@transaction.atomic
def create_products_from_import(*, rows, actor, batch_id, source_filename):
    if not _is_manager(actor):
        raise ValidationError({"actor": "Only an active owner or manager can import products."})
    normalized = _normalize_confirmation_rows(rows)
    filename = _safe_filename(source_filename)
    products = []
    try:
        with transaction.atomic():
            for row in normalized:
                product = Product.objects.create(
                    sku=row["sku"],
                    name=row["name"],
                    cost_price=Decimal(row["cost_price"]),
                    selling_price=Decimal(row["selling_price"]),
                    on_hand=Decimal("0.000"),
                    reorder_level=Decimal(row["reorder_level"]),
                    active=True,
                )
                opening_stock = Decimal(row["opening_stock"])
                if opening_stock > 0:
                    receive_stock(
                        product=product,
                        quantity=opening_stock,
                        unit_cost=Decimal(row["cost_price"]),
                        reason=f"Opening stock imported from {filename} (batch {batch_id}).",
                        actor=actor,
                    )
                products.append(product)
            record_audit_event(
                event_type=AuditEvent.EventType.PRODUCT_IMPORT,
                object_type="product_import",
                object_id=str(batch_id),
                actor=actor,
                reason=f"Imported {len(products)} products from {filename}.",
                after_state={"product_ids": [product.pk for product in products], "count": len(products)},
                metadata={"source_filename": filename, "skus": [product.sku for product in products]},
            )
    except IntegrityError as exc:
        raise ValidationError("An SKU was reserved after preview. No products were imported; upload and preview the file again.") from exc
    return products
